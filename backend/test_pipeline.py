"""
阶段1 自测脚本（联调用，可重复运行）
============================================
作用：不依赖前端，直接验证「解析 → 切分 → 向量化入库 → 列表」链路。
运行方式（在 backend 目录下）：
  python test_pipeline.py
说明：首次入库时 Chroma 会自动下载本地离线 embedding 模型（约 80MB），
      需要联网一次，之后离线可用。
============================================
"""

import io
import uuid
from datetime import datetime

import docx  # python-docx

import parser
import vectorstore


def make_test_docx_bytes() -> bytes:
    """在内存里造一份带标题的测试 Word，返回字节流。"""
    d = docx.Document()
    d.add_heading('第一章 创业担保贷款政策', level=1)
    d.add_paragraph(
        '创业担保贷款对象包括在自治区内依法注册登记的小微企业、个体工商户以及'
        '符合条件的个人创业者。小微企业申请应当经营正常、无重大不良信用记录，'
        '吸纳就业人数达到规定要求。'
    )
    d.add_heading('第二章 一次性创业补贴', level=1)
    d.add_paragraph(
        '对首次创办小微企业且正常经营满一年以上的创业者，给予一次性创业补贴 5000 元。'
        '对吸纳重点群体就业达到一定比例的小微企业，按规定落实社会保险补贴。'
    )
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def main():
    # 1) 造测试文档并解析切分
    data = make_test_docx_bytes()
    source_type, paragraphs = parser.parse_document(data, '测试政策.docx')
    chunks = parser.chunk_paragraphs(paragraphs)

    print('来源类型:', source_type, '| 段落数:', len(paragraphs), '| 切出片段数:', len(chunks))
    for c in chunks:
        print('--- chunk', c['chunk_index'],
              '| page=', c['page'], '| chapter=', c['chapter'])
        print('   ', c['text'][:80].replace('\n', ' '), '...')

    # 2) 向量化入库
    doc_id = uuid.uuid4().hex[:12]
    uploaded_at = datetime.now().isoformat(timespec='seconds')
    n = vectorstore.add_chunks(doc_id, '测试政策.docx', source_type, chunks, uploaded_at)
    print('\n入库片段数:', n)

    # 3) 列表聚合结果
    print('\n当前文档列表:')
    for d in vectorstore.list_documents():
        print('  ', d)


if __name__ == '__main__':
    main()
